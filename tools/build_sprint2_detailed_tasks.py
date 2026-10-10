import ast
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "tools" / "build_sprint2_doc.py"
OUT = ROOT / "Danh_sach_cong_viec_chi_tiet_Sprint_2.docx"
BLUE, LIGHT, PALE, GRAY, WHITE = "1F4E78", "D9EAF7", "F4F8FC", "D9D9D9", "FFFFFF"

def extract_backlog():
    tree = ast.parse(SOURCE.read_text(encoding="utf-8"))
    for node in tree.body:
        if isinstance(node, ast.Assign) and any(isinstance(t, ast.Name) and t.id == "backlog_groups" for t in node.targets):
            return ast.literal_eval(node.value)
    raise RuntimeError("Không tìm thấy backlog_groups")

def font(run, bold=None, color="000000"):
    run.font.name = "Times New Roman"
    rpr = run._element.get_or_add_rPr()
    for tag in ("w:ascii", "w:hAnsi", "w:eastAsia"):
        rpr.rFonts.set(qn(tag), "Times New Roman")
    run.font.size = Pt(13)
    run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None: run.bold = bold

def shade(cell, fill):
    tcpr = cell._tc.get_or_add_tcPr(); el = tcpr.find(qn("w:shd"))
    if el is None: el = OxmlElement("w:shd"); tcpr.append(el)
    el.set(qn("w:fill"), fill)

def border(cell):
    tcpr = cell._tc.get_or_add_tcPr(); box = tcpr.find(qn("w:tcBorders"))
    if box is None: box = OxmlElement("w:tcBorders"); tcpr.append(box)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}"); el.set(qn("w:val"), "single"); el.set(qn("w:sz"), "6"); el.set(qn("w:color"), GRAY); box.append(el)

def margins(cell):
    tcpr = cell._tc.get_or_add_tcPr(); mar = OxmlElement("w:tcMar")
    for edge, val in (("top",100),("start",120),("bottom",100),("end",120)):
        el=OxmlElement(f"w:{edge}"); el.set(qn("w:w"),str(val)); el.set(qn("w:type"),"dxa"); mar.append(el)
    tcpr.append(mar)

def set_cell(cell, text, bold=False, color="000000", align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.text=""; p=cell.paragraphs[0]; p.alignment=align; p.paragraph_format.space_after=Pt(0); p.paragraph_format.line_spacing=1.05
    font(p.add_run(str(text)), bold, color); cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER; border(cell); margins(cell)

def table(doc, headers, rows, widths):
    t=doc.add_table(rows=1, cols=len(headers)); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False
    for i,h in enumerate(headers): set_cell(t.rows[0].cells[i],h,True,WHITE,WD_ALIGN_PARAGRAPH.CENTER); shade(t.rows[0].cells[i],BLUE); t.rows[0].cells[i].width=Inches(widths[i])
    trpr=t.rows[0]._tr.get_or_add_trPr(); rep=OxmlElement("w:tblHeader"); rep.set(qn("w:val"),"true"); trpr.append(rep)
    for ri,row in enumerate(rows):
        cells=t.add_row().cells
        for i,val in enumerate(row): set_cell(cells[i],val,False); cells[i].width=Inches(widths[i]); shade(cells[i], PALE if ri%2 else WHITE)
    return t

def para(doc, text="", bold=False, before=0, after=5, align=None, keep=False):
    p=doc.add_paragraph(); p.paragraph_format.space_before=Pt(before); p.paragraph_format.space_after=Pt(after); p.paragraph_format.line_spacing=1.15; p.paragraph_format.keep_with_next=keep
    if align is not None: p.alignment=align
    font(p.add_run(text),bold); return p

def heading(doc,text,level=1):
    p=doc.add_paragraph(style=f"Heading {level}"); p.paragraph_format.space_before=Pt(10); p.paragraph_format.space_after=Pt(5); p.paragraph_format.keep_with_next=True; font(p.add_run(text),True); return p

def bullets(doc, items, numbered=False):
    for idx,item in enumerate(items,1):
        p=doc.add_paragraph(style="List Bullet" if not numbered else None)
        p.paragraph_format.left_indent=Inches(.25)
        p.paragraph_format.first_line_indent=Inches(-.18)
        p.paragraph_format.space_after=Pt(3); p.paragraph_format.line_spacing=1.1
        font(p.add_run((f"{idx}. " if numbered else "") + item))

story_context = {
"Chung":"Hạ tầng dùng chung cho tất cả chức năng Sprint 2. Nếu phần này sai, các module phía sau dễ mất dữ liệu, sai quyền hoặc khó kiểm tra lỗi.",
"US01":"Ứng viên tự tạo tài khoản, xác thực email, chọn chương trình, lưu nháp và nộp một hồ sơ cho mỗi chương trình.",
"US02":"HR xem hồ sơ đã nộp, lọc, mở chi tiết, duyệt hoặc từ chối; hồ sơ nháp không được xuất hiện.",
"US03":"Kết quả xét duyệt được xếp vào hàng đợi và gửi email theo giờ cố định; HR có thể sửa hoặc thu hồi trước giờ chốt.",
"US04":"HR quản lý hợp đồng của người đã được nhận, tải PDF và giữ lịch sử phiên bản.",
"US05":"Thực tập sinh đọc đúng phiên bản hợp đồng hiện hành rồi xác nhận hoặc từ chối xác nhận.",
"US06":"HR tạo chương trình theo phòng ban, đặt thời hạn nhận hồ sơ, ngày bắt đầu/kết thúc, chỉ tiêu, thứ làm việc và yêu cầu giấy tờ.",
"US07":"HR chia thực tập sinh đã được nhận cho mentor cùng phòng ban và có thể đổi mentor sau đó.",
"US08":"Thực tập sinh xem lịch cá nhân; mốc bắt đầu/kết thúc ưu tiên ngày hợp đồng, nếu chưa có thì lấy từ chương trình.",
"US09":"Thực tập sinh check-in/check-out bằng giờ máy chủ, xem lịch sử và sửa trong thời hạn cho phép; HR sửa khi quá hạn.",
"US10":"HR xem báo cáo chuyên cần từ dữ liệu chấm công, thứ làm việc và ngày nghỉ chung; nghỉ phép hiện ghi Chưa có dữ liệu.",
}

module_info = [
(range(1,7),"Nền tảng dữ liệu và dùng chung","src/main/resources, package util/service/filter, migration SQL"),
(range(7,18),"Đăng ký và nộp hồ sơ","User, application, document; servlet đăng ký/xác thực/hồ sơ; JSP ứng viên"),
(range(18,24),"Chương trình thực tập","internship_programs, ProgramDAO/Service/Servlet, JSP chương trình"),
(range(24,35),"Xét duyệt và email kết quả","application_decisions, email_result_queue, servlet HR và scheduler"),
(range(35,42),"Hợp đồng","contracts, contract_versions, ContractDAO/Service/Servlet và JSP hai phía"),
(range(42,51),"Mentor và lịch cá nhân","mentor_assignment_history, calendar_events, service và JSP"),
(range(51,63),"Chấm công và báo cáo","attendance_records, logs, holidays, report service và JSP"),
(range(63,68),"Hoàn thiện và nghiệm thu","menu, bảo mật, test, cấu hình triển khai và tài liệu"),
]

def num(task_id): return int(task_id.split("-")[1])
def module_for(n):
    for rng,name,areas in module_info:
        if n in rng: return name,areas
    return "Khác",""

def role_steps(owner):
    if "DB" in owner:
        return ["Vẽ hoặc ghi rõ bảng, cột, kiểu dữ liệu, khóa chính, khóa ngoại, unique và index trước khi viết SQL.","Viết migration có thể chạy trên cơ sở dữ liệu đang có; kiểm tra dữ liệu cũ trước khi thêm NOT NULL hoặc constraint.","Cập nhật câu lệnh seed/demo nếu chức năng cần dữ liệu mẫu và ghi lại cách rollback."]
    if owner == "FE" or "FE" in owner:
        return ["Đối chiếu wireframe và trạng thái nghiệp vụ; liệt kê trạng thái bình thường, trống, lỗi, đang tải và không có quyền.","Tạo hoặc sửa JSP trong WEB-INF/views; dùng CSS hiện có, không đưa quy tắc nghiệp vụ chính vào JavaScript.","Kết nối form với URL servlet, giữ dữ liệu người dùng khi validation lỗi và escape dữ liệu hiển thị."]
    if owner == "QA":
        return ["Chuyển tiêu chí chấp nhận thành test case gồm dữ liệu chuẩn bị, bước thao tác và kết quả mong đợi.","Kiểm tra cả luồng đúng, dữ liệu biên, quyền truy cập, bấm lặp, hai tab và lỗi kết nối.","Ghi bằng chứng test, mã lỗi và dữ liệu tái hiện; chạy regression các chức năng liên quan sau khi sửa."]
    if owner == "DevOps":
        return ["Liệt kê biến môi trường và thư mục cần cấu hình; không để mật khẩu hoặc token trong source.","Chuẩn bị cấu hình cho môi trường local và triển khai, kèm giá trị mẫu an toàn.","Kiểm tra khởi động lại ứng dụng, quyền thư mục, log lỗi và khả năng phục hồi job đang chạy."]
    return ["Xác định request, response, trạng thái và lỗi mà servlet/service phải xử lý.","Tạo hoặc sửa model, DAO và service trước; servlet chỉ nhận dữ liệu, gọi service và chuyển sang JSP hoặc trả lỗi.","Bọc các thao tác nhiều bước trong transaction, kiểm tra quyền và dữ liệu lại ở phía server.","Viết test cho quy tắc cốt lõi và trường hợp request bị gửi lặp."]

specific = {
1:["Tạo file migration-002-sprint2.sql thay vì sửa lịch sử migration đã chạy.","Thứ tự tạo bảng phải tôn trọng khóa ngoại; chạy thử trên schema rỗng và một bản sao schema hiện tại.","Ghi lệnh rollback hoặc kế hoạch khôi phục cho từng nhóm bảng."],
2:["Tạo lớp hằng số hoặc enum Java cho từng vòng đời.","Tên DB và Java phải có bảng ánh xạ rõ, ví dụ SUBMITTED không được nơi khác gọi là NEW.","Mọi chuyển trạng thái phải đi qua service và từ chối chuyển sai."],
3:["Tạo AuditService nhận userId, actionCode, tableName, recordId, oldValues, newValues và IP.","Chuyển object cũ/mới thành JSON; không ghi mật khẩu, token hay nội dung tệp.","Gọi service trong cùng transaction với thao tác chính khi cần tính nhất quán."],
4:["Thêm permission_code cho từng module và INSERT IGNORE vào migration.","Cập nhật AuthorizationFilter hoặc ánh xạ URL để kiểm quyền server-side.","Tạo bảng ma trận role - permission để QA dùng kiểm thử."],
5:["Lưu tệp ngoài thư mục public; sinh tên vật lý ngẫu nhiên nhưng giữ tên gốc trong DB.","Chặn đuôi và MIME không hợp lệ, giới hạn dung lượng, chống ../ trong tên tệp.","Tạo endpoint download kiểm tra người dùng có quyền với bản ghi trước khi trả file."],
6:["Chuẩn bị ít nhất một Admin, HR, Mentor, Candidate và Intern.","Tạo chương trình sắp mở, đang mở, đã đóng; hồ sơ nháp, đã nộp, đã duyệt, từ chối.","Tạo hợp đồng, phân công, chấm công đủ trường hợp để test báo cáo."],
7:["Thêm email_verified_at và bảng/token xác thực có token_hash, expires_at, used_at, resend_count.","Không lưu token thuần; unique email vẫn được giữ.","Tạo index phục vụ tìm token còn hiệu lực."],
8:["GET hiển thị form, POST validate họ tên/email/mật khẩu/xác nhận.","Kiểm email trùng; băm mật khẩu bằng cơ chế hiện có; tạo role CANDIDATE.","Tạo token xác thực trong cùng transaction rồi yêu cầu EmailService gửi."],
9:["Endpoint verify nhận token, băm và tìm token chưa dùng/còn hạn.","Gửi lại phải vô hiệu token cũ, tạo token mới và giới hạn số lần/thời gian.","Thông báo không được tiết lộ email có tồn tại khi có nguy cơ dò tài khoản."],
10:["Tạo register.jsp, verify-result.jsp và resend-verification.jsp đồng bộ style login.","Hiển thị yêu cầu mật khẩu trước khi gửi và lỗi ngay dưới trường.","Sau thành công hướng dẫn người dùng kiểm tra email và cách gửi lại."],
11:["Tạo bảng applications gắn user_id và program_id, unique hai cột.","Thêm university, major, study_year, GPA, status, submitted_at, version.","Không dùng intern_profiles cho hồ sơ ứng tuyển vì chỉ tạo hồ sơ TTS chính thức sau khi được nhận."],
12:["Tạo danh mục yêu cầu giấy tờ theo chương trình và bảng tệp của hồ sơ.","Mỗi yêu cầu có loại, tên hiển thị, bắt buộc/tùy chọn, định dạng và dung lượng.","Tạo unique tránh một hồ sơ có nhiều tệp cùng loại nếu nghiệp vụ không cho phép."],
13:["Chỉ trả chương trình có trạng thái Đang nhận hồ sơ cho thao tác bắt đầu.","Khi người dùng chọn chương trình, lấy hồ sơ đã có hoặc tạo một bản DRAFT duy nhất.","Lưu nháp từng lần submit và trả lại đầy đủ dữ liệu khi mở lại."],
14:["Upload từng tệp độc lập để lỗi một tệp không xóa dữ liệu form.","Chỉ cho thay/xóa tệp khi hồ sơ DRAFT hoặc NEEDS_SUPPLEMENT.","Khi thay tệp, cập nhật metadata và xử lý tệp cũ theo chính sách lưu trữ."],
15:["Khóa hoặc kiểm version hồ sơ trong transaction.","Kiểm email đã xác thực, chương trình đang mở, trường bắt buộc, CV và giấy tờ bắt buộc.","Cập nhật SUBMITTED và submitted_at đúng một lần; request lặp trả kết quả hiện tại."],
16:["Tạo trang danh sách/chỉ tiết chương trình, form hồ sơ theo bước và trang Hồ sơ của tôi.","Nút Lưu nháp luôn rõ; nút Nộp chỉ mở khi có thể nộp.","Hiển thị chính xác giấy tờ bắt buộc/tùy chọn và lý do hồ sơ bị khóa."],
17:["Test email trùng, token hết hạn/đã dùng, gửi lại quá nhanh.","Test chương trình đóng khi đang soạn, thiếu giấy tờ, file sai, bấm Nộp hai lần.","Dùng hai tài khoản thử đổi ID hồ sơ để xác nhận không xem được dữ liệu nhau."],
18:["Bổ sung riêng ngày mở/đóng hồ sơ và ngày bắt đầu/kết thúc thực tập.","Lưu work_days theo cấu trúc nhất quán, mặc định T2-T6 và ít nhất một ngày.","Thêm version cho optimistic locking và closed_early_at cho đóng sớm."],
19:["Tính trạng thái từ ngày hiện tại và cờ đóng sớm thay vì tin giá trị do form gửi.","Kiểm tên trùng trong cùng phòng ban và kiểm thứ tự bốn mốc ngày.","Khi update dùng WHERE id=? AND version=?; nếu 0 dòng thì báo dữ liệu đã đổi."],
20:["Tạo URL list/create/edit/detail/close/delete với kiểm quyền HR.","Xóa chỉ khi chưa có hồ sơ; đóng sớm không xóa dữ liệu.","Khi đã có hồ sơ nộp, bỏ qua thay đổi yêu cầu giấy tờ nhưng vẫn lưu trường hợp lệ khác theo nghiệp vụ."],
21:["Viết truy vấn đếm tổng hồ sơ, đã duyệt và được nhận theo program_id.","Không đếm hồ sơ nháp; thống nhất tập trạng thái được tính.","Từ chi tiết tạo link sang danh sách hồ sơ với programId đã điền."],
22:["Form có tên, phòng ban, chỉ tiêu, bốn mốc ngày, thứ làm việc, giấy tờ và mô tả.","Đặt ngày bắt đầu/kết thúc ngay trong form chương trình, không tạo màn hình riêng.","Trang chi tiết hiển thị thống kê, cảnh báo vượt chỉ tiêu và nút đóng/sửa/xóa phù hợp."],
23:["Test đúng ngày mở/đóng, ngày không hợp lệ, tên trùng và bỏ hết thứ làm việc.","Mở hai trình duyệt sửa cùng chương trình để test version.","Test đóng sớm, gia hạn mở lại, khóa giấy tờ và xóa chương trình có hồ sơ."],
24:["Tạo application_decisions lưu application_id, decision, reason, decided_by, decided_at.","Giữ lịch sử thay vì chỉ ghi đè một cột status.","Thêm version hoặc khóa phục vụ hai HR xử lý đồng thời."],
25:["DAO nhận đối tượng filter và dùng PreparedStatement cho mọi điều kiện.","Mặc định chỉ lấy SUBMITTED/UNDER_REVIEW; hỗ trợ phân trang và đếm tổng.","Thêm index cho program, status, submitted_at và các trường lọc thường dùng."],
26:["Khi HR mở lần đầu, chuyển SUBMITTED sang UNDER_REVIEW bằng update có điều kiện.","Chi tiết gồm thông tin cấu trúc và link tải tệp qua endpoint kiểm quyền.","Không thay đổi hồ sơ đã có quyết định."],
27:["Nhận danh sách applicationId và hành động; tải lại trạng thái từng hồ sơ lúc xác nhận.","Duyệt không cần lý do; từ chối bắt buộc lý do riêng từng hồ sơ.","Hồ sơ bị HR khác xử lý được bỏ qua, còn các hồ sơ khác vẫn tiếp tục; trả báo cáo theo từng ID."],
28:["Tính quota sau khi cộng cả lô đang duyệt.","Nếu đạt/vượt, trả thông tin hiện tại, sau thao tác và chỉ tiêu để giao diện xác nhận lại.","Dùng confirm token hoặc tham số xác nhận để tránh tự động vượt quota."],
29:["Bộ lọc phải giữ giá trị sau submit và có nút xóa lọc.","Chọn tất cả chỉ áp dụng tập kết quả lọc và hiển thị rõ số lượng.","Modal từ chối có lý do theo từng dòng; đánh dấu đúng dòng thiếu."],
30:["Tạo hàng đợi có application_id, result_type, payload, status, scheduled_at, attempts, locked_at, sent_at, error.","Đảm bảo tối đa một bản PENDING cho mỗi hồ sơ bằng constraint hoặc transaction.","Thêm khóa idempotency để một thông báo không phát sinh hai email."],
31:["Tạo decision và queue item trong cùng transaction.","Khi sửa quyết định, hủy item cũ rồi tạo item mới.","Payload chỉ chứa dữ liệu ổn định cần thiết; trước gửi vẫn kiểm trạng thái mới nhất."],
32:["Đọc giờ gửi từ cấu hình; lấy lô PENDING đến hạn và đánh dấu đang xử lý bằng khóa.","Thành công cập nhật SENT và sent_at; lỗi tạm thời tăng attempts và hẹn retry.","Lỗi vĩnh viễn hoặc quá số lần chuyển FAILED; khởi động lại không gửi lại item SENT."],
33:["Danh sách queue có lọc trạng thái, loại kết quả, ứng viên và thời gian.","Thu hồi chỉ với PENDING trước khi bị job chốt; đưa hồ sơ về UNDER_REVIEW và audit.","Sửa kết quả dẫn về màn xét duyệt và tạo queue item mới sau xác nhận."],
34:["Chạy hai request HR cùng lúc trên cùng hồ sơ và kiểm chỉ một quyết định thắng.","Dừng job giữa lô rồi chạy lại để kiểm không gửi trùng.","Test sửa nhiều lần, thu hồi sát giờ, lỗi SMTP và việc ẩn kết quả trước SENT."],
35:["Bổ sung confirmation_deadline, uploaded_by, reject_reason, version/checksum.","Tạo contract_versions giữ file_url, version_no, checksum và trạng thái hiệu lực.","Ràng buộc một hợp đồng hiện hành cho mỗi TTS theo quy tắc nghiệp vụ."],
36:["Chỉ lấy người có kết quả cuối cùng Được nhận.","Form mặc định ngày theo chương trình; hiển thị mentor hoặc Chưa có mentor.","Sau upload tạo contract và version 1, trạng thái chờ xác nhận."],
37:["Thay tệp tạo version mới, không ghi đè file cũ; version cũ mất hiệu lực.","Không cho thay tùy tiện sau khi TTS đã xác nhận nếu nghiệp vụ không cho phép.","Xóa hợp đồng phải audit và làm lịch/chấm công quay về ngày chương trình."],
38:["Danh sách có TTS, chương trình, mentor, thời gian, phiên bản và trạng thái xác nhận.","Form kiểm PDF, số hợp đồng, ngày, phụ cấp và hạn xác nhận.","Trước gửi hiển thị tóm tắt để HR kiểm tra."],
39:["Tìm hợp đồng hiện hành theo user đăng nhập, không nhận internId tùy ý.","Download PDF qua endpoint kiểm quyền và checksum/version hiện hành.","Xác nhận hoặc từ chối bằng update version; lưu thời gian, IP và lý do từ chối."],
40:["Hiển thị số hợp đồng, thời gian, phụ cấp, hạn, phiên bản và link PDF.","Yêu cầu checkbox đã đọc trước khi bật nút Xác nhận.","Nếu version đổi hoặc hết hạn, khóa thao tác và yêu cầu tải lại."],
41:["Test thay file ngay trước lúc xác nhận, bấm hai tab và hợp đồng hết hạn.","Thử tải hợp đồng người khác bằng đổi URL.","Sau sửa/xóa hợp đồng, kiểm lịch và giới hạn chấm công lấy đúng nguồn ngày."],
42:["Tạo lịch sử intern_profile_id, old_mentor_id, new_mentor_id, assigned_by, assigned_at.","mentor_id trên intern_profiles là trạng thái hiện hành; history là dữ liệu truy vết.","Mọi đổi mentor cập nhật cả hai trong một transaction."],
43:["TTS đủ điều kiện là người đã được nhận trong chương trình.","Mentor đủ điều kiện phải ACTIVE và department_id trùng phòng ban chương trình.","Đếm tải hiện tại từ phân công hiện hành, không chỉ nhóm đang lọc."],
44:["Sắp mentor theo số TTS tăng dần rồi theo mentor_code/id cố định.","Lần lượt gán TTS chưa có mentor cho người ít tải nhất và cập nhật tải giả lập.","Chỉ trả đề xuất về giao diện; chưa ghi DB trước khi HR xác nhận."],
45:["Cho lọc theo chương trình/phòng ban/trạng thái phân công.","Hiển thị mentor hiện tại, mentor đề xuất và tải trước/sau.","Khi xác nhận gửi toàn bộ lựa chọn, server kiểm lại mentor còn hợp lệ."],
46:["Test phòng ban không mentor, mentor inactive và mentor khác phòng ban.","Hai HR cùng xác nhận phân công để kiểm version/transaction.","Kiểm mentor mới xem được lịch sử, mentor cũ mất quyền sau đổi."],
47:["Tạo event có tiêu đề, mô tả, start/end, all_day, creator, status.","Tạo recipient theo INTERN, MENTOR_GROUP hoặc PROGRAM.","Sự kiện CANCELLED không xuất hiện nhưng vẫn giữ để audit."],
48:["Xác định TTS từ user đăng nhập; từ chối nếu chưa được nhận.","Tạo hai mốc bắt đầu/kết thúc động từ hợp đồng hiện hành, nếu chưa có lấy chương trình.","Lấy sự kiện gửi riêng, theo mentor hiện tại hoặc chương trình; lọc giao với khoảng ngày."],
49:["Hỗ trợ tháng/tuần, trước/sau, Hôm nay và đánh dấu ngày hiện tại.","Ngày nhiều mục hiển thị +N mục khác; sự kiện nhiều ngày xuất hiện trên các ngày bao phủ.","Chi tiết nêu nguồn mốc hoặc người tạo sự kiện; giao diện chỉ đọc."],
50:["Test chưa được nhận, được nhận nhưng chưa có sự kiện và chưa có hợp đồng.","Test sự kiện hủy/nhiều ngày/qua tháng và đổi mentor.","Đổi URL hoặc request để lấy lịch người khác phải bị từ chối."],
51:["Mở rộng trạng thái WORKING, COMPLETED, MISSING_CHECKOUT, CANCELLED.","Thêm total_minutes, edited flag; tạo attendance_change_logs giữ before/after, reason, editor.","Cấu hình số ngày tự sửa và index intern_profile_id + work_date."],
52:["Tìm kỳ thực tập từ hợp đồng hiện hành, nếu chưa có lấy chương trình.","Check-in insert theo ngày bằng giờ server; unique ngăn lần hai.","Check-out update bản ghi WORKING, giờ ra phải sau giờ vào và cùng ngày; tính total_minutes."],
53:["Job tìm bản ghi WORKING của ngày trước và chuyển MISSING_CHECKOUT.","Không điền check_out_time và không tính tổng giờ.","Update có điều kiện để job chạy lại không thay đổi bản ghi đã xử lý."],
54:["TTS chỉ sửa trong hạn; HR sửa quá hạn; mọi sửa bắt buộc lý do.","Validate giờ không tương lai, cùng ngày, checkout sau checkin.","Hủy không xóa bản ghi; không tính giờ và cho phép tạo lượt hợp lệ cho ngày đó theo thiết kế constraint."],
55:["Lịch sử tháng chỉ trả dữ liệu đúng quyền; tính tổng bỏ CANCELLED và thiếu checkout.","Mentor chỉ xem TTS đang được mình hướng dẫn; HR lọc theo chương trình/TTS/khoảng ngày.","Hiển thị change log và giá trị gốc cho bản ghi đã sửa."],
56:["Khối hôm nay hiển thị giờ server, kỳ thực tập và nút phù hợp trạng thái.","Bảng tháng có tổng giờ, trạng thái, nhãn chỉnh sửa và hành động bổ sung/sửa.","Form sửa giải thích hạn, bắt buộc lý do và giữ dữ liệu khi validation lỗi."],
57:["Gửi check-in/check-out đồng thời từ hai tab và kiểm chỉ một thao tác thành công.","Test quên checkout qua ngày, thêm ngày bị quên, giờ sai và ngoài kỳ.","Test quyền TTS/mentor/HR và tác động khi đổi mentor."],
58:["Tạo public_holidays với holiday_date unique, name và audit.","work_days lấy từ chương trình; dữ liệu cũ thiếu thì mặc định T2-T6 và báo HR.","Thêm index attendance phục vụ tổng hợp theo người và ngày."],
59:["Cắt khoảng lọc theo kỳ thực tập của từng người.","Sinh ngày kỳ vọng từ work_days, loại holiday, chỉ tính ngày đã qua cho vắng.","Ngày có bản ghi không hủy là đi làm; thiếu checkout đếm riêng; ngoài ngày kỳ vọng chỉ cộng giờ.","Tỷ lệ = ngày đi làm / ngày kỳ vọng đã qua; mẫu số 0 trả null để hiển thị gạch ngang."],
60:["Filter theo chương trình, phòng ban, tên/email và khoảng ngày.","Tính tổng trên toàn bộ tập lọc trước khi phân trang.","Chi tiết ngày phân loại đi làm, vắng, nghỉ chung, ngoài kỳ vọng, thiếu checkout hoặc hủy."],
61:["Bảng tổng hợp có các cột đã chốt và cột Nghỉ phép = Chưa có dữ liệu.","Cho sắp xếp, phân trang và mở chi tiết từng ngày; báo cáo không sửa trực tiếp chấm công.","Màn ngày nghỉ chung hỗ trợ thêm/sửa/xóa, xác nhận khi thay đổi làm tính lại quá khứ."],
62:["Test TTS bắt đầu giữa tháng, kết thúc giữa tháng và khoảng chỉ có ngày nghỉ.","Test khai báo lễ sau khi đã chấm công, cuối tuần có chấm công và hôm nay chưa kết thúc.","Kiểm tổng nhóm không thay đổi theo trang và dữ liệu cũ thiếu work_days."],
63:["Thêm menu theo role, không chỉ ẩn bằng CSS mà URL vẫn phải được filter bảo vệ.","Dùng chung badge cho các trạng thái và thông điệp dễ hiểu.","Mỗi màn có trạng thái trống, lỗi tải và nút thử lại phù hợp."],
64:["Bổ sung CSRF token cho form ghi dữ liệu và escape output JSP.","Rà toàn bộ query dùng PreparedStatement, endpoint kiểm ownership để chống IDOR.","Test file giả MIME, tên tệp nguy hiểm, request lớn và rate limit email."],
65:["Viết kịch bản end-to-end: tạo chương trình -> đăng ký -> nộp -> duyệt -> email -> hợp đồng -> mentor -> chấm công -> báo cáo.","Cố ý gây lỗi giữa transaction để kiểm rollback.","Chạy regression đăng nhập, quên/đổi mật khẩu và chức năng Sprint 1."],
66:["Đưa SMTP host/port/user/password, send time, upload path và retry count ra biến môi trường.","Tạo thư mục upload với quyền tối thiểu và chính sách backup/retention.","Hướng dẫn cách xem log job, xử lý FAILED và khởi động lại an toàn."],
67:["Chuẩn bị dữ liệu demo xuyên suốt và phân công người trình bày từng lát cắt.","Đối chiếu từng tiêu chí chấp nhận; lỗi chưa xong phải có owner và thời hạn.","Chốt tài liệu cài đặt, migration, tài khoản demo và biên bản nghiệm thu."],
}

def deliverables(n, owner, areas):
    items=[]
    if "DB" in owner: items.append("Migration SQL và sơ đồ/bảng mô tả thay đổi dữ liệu")
    if owner=="FE" or "FE" in owner: items.append("JSP/CSS/JavaScript của màn hình, gồm trạng thái lỗi và trống")
    if owner not in ("FE","QA","DevOps") or "BE" in owner: items.append("Model/DAO/Service/Servlet hoặc job cần thiết, kèm xử lý lỗi")
    if owner=="QA": items.append("Test case, dữ liệu test và biên bản kết quả")
    if owner=="DevOps": items.append("Tệp cấu hình mẫu và hướng dẫn triển khai/vận hành")
    items.append("Bằng chứng tự kiểm tra và ghi chú cho người review")
    return items

def verification(n, acceptance):
    base=[acceptance,"Không làm hỏng dữ liệu hoặc chức năng đã có liên quan.","Trường hợp không có quyền hoặc dữ liệu không hợp lệ được server từ chối rõ ràng."]
    if n in range(7,63): base.append("Kiểm tra request gửi lặp hoặc hai tab không tạo dữ liệu trùng.")
    return base

backlog=extract_backlog()
doc=Document(); sec=doc.sections[0]; sec.page_width=Inches(8.5); sec.page_height=Inches(11); sec.top_margin=sec.bottom_margin=Inches(.7); sec.left_margin=sec.right_margin=Inches(.75)
for name in ("Normal","Title","Heading 1","Heading 2","Heading 3","List Bullet","List Number"):
    st=doc.styles[name]; st.font.name="Times New Roman"; st.font.size=Pt(13); st.font.color.rgb=RGBColor(0,0,0)
    for tag in ("w:ascii","w:hAnsi","w:eastAsia"): st._element.get_or_add_rPr().rFonts.set(qn(tag),"Times New Roman")

t=doc.add_paragraph(style="Title"); t.alignment=WD_ALIGN_PARAGRAPH.CENTER; font(t.add_run("DANH SÁCH CÔNG VIỆC CHI TIẾT SPRINT 2"),True)
para(doc,"Hệ thống quản lý thực tập sinh",True,after=8,align=WD_ALIGN_PARAGRAPH.CENTER)
para(doc,"Tài liệu giao việc dành cho thành viên chưa nắm đầy đủ nghiệp vụ và cấu trúc hệ thống",False,after=14,align=WD_ALIGN_PARAGRAPH.CENTER)
table(doc,["Nội dung","Cách sử dụng"],[
["Mỗi công việc","Đọc mục tiêu, bối cảnh, nơi cần làm, từng bước, kết quả bàn giao và cách kiểm tra trước khi bắt đầu code."],
["Mã S2","Giữ nguyên mã S2-01 đến S2-67 để đối chiếu tài liệu backlog và đưa vào Jira hoặc Trello."],
["Phân công","Một người là chủ sở hữu chính; người khác có thể review hoặc hỗ trợ nhưng không để task không có người chịu trách nhiệm."],
["Quy tắc hoàn thành","Chỉ chuyển Done khi đã có sản phẩm bàn giao, tự kiểm tra và người review xác nhận."],
],[1.5,5.2])

heading(doc,"1 Cách đọc tài liệu",1)
bullets(doc,[
"DB là công việc thiết kế dữ liệu và migration. BE là Java model, DAO, service, servlet, filter hoặc job. FE là JSP, CSS và JavaScript. QA là kiểm thử. DevOps là cấu hình chạy và triển khai.",
"Phần Nơi thực hiện là vùng dự kiến trong repository, không bắt buộc đúng tên lớp. Thành viên phải kiểm tra code hiện tại trước khi tạo file mới để tránh trùng chức năng.",
"Phần Phụ thuộc nghĩa là task trước phải đủ ổn định hoặc đã thống nhất hợp đồng dữ liệu. Có thể code song song bằng mock nhưng không được tích hợp khi hợp đồng chưa chốt.",
"Mọi thao tác ghi dữ liệu quan trọng phải kiểm quyền phía server, có transaction khi nhiều bước, chống gửi lặp và ghi audit khi tài liệu nghiệp vụ yêu cầu.",
])
heading(doc,"2 Thứ tự làm việc bắt buộc cho mỗi thành viên",1)
bullets(doc,[
"Bước 1: đọc user story liên quan trong thư mục Nghiệp vụ và xác nhận đầu vào, trạng thái, ngoại lệ.",
"Bước 2: đọc model, DAO, servlet, JSP và schema hiện tại liên quan; ghi rõ sẽ sửa file nào và tạo file nào.",
"Bước 3: thống nhất dữ liệu và URL với người làm phần còn lại trước khi code.",
"Bước 4: code theo lớp dữ liệu -> service -> servlet -> giao diện; không đặt quy tắc nghiệp vụ trong JSP.",
"Bước 5: tự test luồng đúng, luồng lỗi, quyền và request lặp; gửi kèm dữ liệu test và ảnh/bằng chứng cho reviewer.",
])
heading(doc,"3 Bảng tra nhanh 67 công việc",1)
quick=[]
for group,rows in backlog:
    for tid,story,owner,summary,dep,acc in rows: quick.append([tid,story,owner,summary,dep])
table(doc,["Mã","Story","Phụ trách","Tên công việc","Phụ thuộc"],quick,[.65,.7,1.0,3.75,1.05])

heading(doc,"4 Hướng dẫn chi tiết từng công việc",1)
for group,rows in backlog:
    doc.add_page_break(); heading(doc,group.replace("4.","4 ",1),1)
    for tid,story,owner,summary,dep,acc in rows:
        n=num(tid); module,areas=module_for(n)
        heading(doc,f"{tid} {summary.rstrip('.')}",2)
        table(doc,["Thuộc phần","Chủ sở hữu gợi ý","Phụ thuộc","Nơi thực hiện dự kiến"],[[story,owner,dep,areas]],[1.0,1.25,1.15,3.8])
        para(doc,"Mục tiêu",True,before=5,after=2,keep=True)
        para(doc,summary.rstrip(".")+". Sau khi hoàn thành, phần việc phải đủ rõ để module khác tích hợp mà không phải đoán trạng thái hoặc cấu trúc dữ liệu.")
        para(doc,"Bối cảnh dễ hiểu",True,before=4,after=2,keep=True)
        contexts=[]
        for key in story.split():
            if key in story_context and story_context[key] not in contexts: contexts.append(story_context[key])
        para(doc," ".join(contexts) if contexts else story_context.get(story,story_context["Chung"]))
        para(doc,"Các bước cần làm",True,before=4,after=2,keep=True)
        steps=[]
        for s in role_steps(owner)+specific.get(n,[]):
            if s not in steps: steps.append(s)
        bullets(doc,steps,numbered=True)
        para(doc,"Kết quả phải bàn giao",True,before=4,after=2,keep=True)
        bullets(doc,deliverables(n,owner,areas))
        para(doc,"Tự kiểm tra trước khi báo hoàn thành",True,before=4,after=2,keep=True)
        bullets(doc,verification(n,acc))
        para(doc,"Khi nào chưa được coi là hoàn thành",True,before=4,after=2,keep=True)
        para(doc,"Chưa được chuyển Done nếu mới dựng giao diện tĩnh, mới tạo bảng nhưng chưa tích hợp, chỉ chạy được dữ liệu mẫu thuận lợi, chưa kiểm quyền, chưa xử lý lỗi hoặc chưa có bằng chứng kiểm thử.",False,after=9)

doc.add_page_break(); heading(doc,"5 Checklist dành cho trưởng nhóm khi giao việc",1)
table(doc,["Đánh dấu","Câu hỏi cần xác nhận"],[["□",x] for x in [
"Task đã có một chủ sở hữu chính và một người review chưa?",
"Người nhận đã đọc user story và biết trạng thái đầu vào/đầu ra chưa?",
"Đã chốt bảng/cột, URL, request/response với task liên quan chưa?",
"Có dữ liệu test đủ luồng đúng, lỗi, quyền và bấm lặp chưa?",
"Có nguy cơ hai người sửa cùng file hoặc cùng bảng không?",
"Task có đầu ra cụ thể để demo và nghiệm thu không?",
"Các lỗi phát hiện đã có mã, owner và thời hạn xử lý chưa?",
]],[.9,5.9])

footer=doc.sections[0].footer; fp=footer.paragraphs[0]; fp.alignment=WD_ALIGN_PARAGRAPH.CENTER; font(fp.add_run("Danh sách công việc chi tiết Sprint 2"))
doc.core_properties.title="Danh sách công việc chi tiết Sprint 2"
doc.core_properties.subject="Hướng dẫn giao việc cho nhóm phát triển hệ thống quản lý thực tập sinh"
doc.core_properties.author="Nhóm phát triển Intern Management"
doc.save(OUT)
print(OUT)
